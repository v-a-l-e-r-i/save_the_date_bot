from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from bot.db import get_all_guests

COLUMNS = [
    ("chat_id", "chat_id"),
    ("username", "username"),
    ("full_name", "Ім'я та прізвище"),
    ("phone", "Телефон"),
    ("company", "Компанія"),
    ("position", "Посада"),
    ("chosen_date", "Обрана дата"),
    ("status", "Статус"),
    ("invited_at", "Дата запрошення"),
    ("first_response_at", "Дата першої відповіді"),
    ("date_history", "Історія зміни дати"),
]


async def build_report(session: AsyncSession, out_path: str) -> str:
    guests = await get_all_guests(session)

    wb = Workbook()
    ws = wb.active
    ws.title = "Guests"
    ws.append([label for _, label in COLUMNS])

    for g in guests:
        history = g.date_history()
        history_str = "; ".join(f"{h['date']} ({h['changed_at'][:16]})" for h in history)
        row = [
            g.chat_id,
            g.username,
            g.full_name,
            g.phone,
            g.company,
            g.position,
            g.chosen_date.isoformat() if g.chosen_date else "",
            g.status,
            g.invited_at.isoformat(timespec="minutes") if g.invited_at else "",
            g.first_response_at.isoformat(timespec="minutes") if g.first_response_at else "",
            history_str,
        ]
        ws.append(row)

    for i, (_, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(label) + 4)

    wb.save(out_path)
    return out_path
